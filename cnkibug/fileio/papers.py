from __future__ import annotations

import csv
import re
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font

from ..cnki.models import Paper, deduplicate_papers, document_type_family, normalize_doi, split_authors


COLUMNS = (
    ("title", "论文标题"), ("authors", "作者"), ("source", "来源"),
    ("publication_date", "发表日期"), ("document_type", "文献类型"), ("doi", "DOI"),
    ("citation_count", "被引次数"), ("download_count", "下载次数"),
    ("paper_keywords", "论文关键词"), ("abstract", "摘要"),
    ("institutions", "作者单位"), ("funds", "基金"), ("classification", "分类号"),
    ("volume", "卷"), ("issue", "期"), ("pages", "页码"),
    ("citation", "引用格式"), ("detail_url", "详情链接"), ("queries", "命中检索项"),
)
DETAIL_FIELDS = {"paper_keywords", "abstract", "institutions", "funds", "classification", "volume", "issue", "pages"}
ALIASES = {label: key for key, label in COLUMNS} | {key: key for key, _ in COLUMNS}
ALIASES.update({"keyword": "queries", "检索项": "queries", "题名": "title"})

_RIS_TYPES = {
    "": "GEN", "journal": "JOUR", "thesis": "THES", "conference": "CONF",
    "book": "BOOK", "book_section": "CHAP", "newspaper": "NEWS", "patent": "PAT",
    "standard": "STAND", "statute": "STAT", "video": "VIDEO",
}


def paper_columns(include_citation: bool = True, include_details: bool = True) -> list[tuple[str, str]]:
    return [(key, label) for key, label in COLUMNS
            if (include_details or key not in DETAIL_FIELDS) and (include_citation or key != "citation")]


def paper_values(paper: Paper, columns: list[tuple[str, str]]) -> list[str]:
    values = []
    for key, _ in columns:
        value = getattr(paper, key)
        if key == "queries":
            value = "\n".join(value)
        elif key == "paper_keywords":
            value = "；".join(split_values(value))
        values.append(ILLEGAL_CHARACTERS_RE.sub("", str(value)))
    return values


def split_values(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"[;；\n]+", value) if part.strip()]


def append_papers(ws, papers: list[Paper], include_citation: bool = True, include_details: bool = True) -> None:
    columns = paper_columns(include_citation, include_details)
    ws.append([label for _, label in columns])
    for paper in papers:
        ws.append([value[:32767] for value in paper_values(paper, columns)])
        for cell in ws[ws.max_row]:
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=False)
        for key in ("detail_url", "doi"):
            column = next((i + 1 for i, (name, _) in enumerate(columns) if name == key), None)
            value = getattr(paper, key)
            if column and value:
                url = f"https://doi.org/{value}" if key == "doi" else value
                if url.startswith(("https://", "http://")):
                    cell = ws.cell(ws.max_row, column)
                    cell.hyperlink = url
                    cell.style = "Hyperlink"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for index, (key, _) in enumerate(columns, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = {
            "title": 55, "abstract": 70, "authors": 24, "source": 26,
            "doi": 32, "detail_url": 35, "queries": 26,
        }.get(key, 18)


def paper_workbook(papers: list[Paper], include_citation: bool = True, include_details: bool = True):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "论文信息"
    append_papers(worksheet, papers, include_citation, include_details)
    return workbook


def write_csv(path: str | Path, papers: list[Paper], include_citation: bool = True, include_details: bool = True) -> None:
    columns = paper_columns(include_citation, include_details)
    with open(path, "w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow([label for _, label in columns])
        writer.writerows(paper_values(paper, columns) for paper in papers)


def _read_rows(rows, query: str = "") -> list[Paper]:
    header = next(rows, ())
    keys = [ALIASES.get(str(value or "").strip(), "") for value in header]
    if "title" not in keys:
        raise ValueError("文件缺少“论文标题”列，请选择 CNKIBug 导出的 XLSX 或 CSV 文件。")
    papers = []
    for row in rows:
        values = {key: str(value).strip() if value is not None else "" for key, value in zip(keys, row) if key}
        if not values.get("title"):
            continue
        raw_queries = values.pop("queries", "")
        paper = Paper(**values)
        paper.queries = [value.strip() for value in raw_queries.splitlines() if value.strip()] or ([query] if query else [])
        paper.doi = normalize_doi(paper.doi)
        papers.append(paper)
    return papers


def read_papers(path: str | Path) -> list[Paper]:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            return deduplicate_papers(_read_rows(iter(csv.reader(stream))))
    if path.suffix.lower() != ".xlsx":
        raise ValueError("请选择 XLSX 或 CSV 文件。")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        papers = []
        for sheet in workbook.worksheets:
            if sheet.max_row == 0:
                continue
            papers.extend(_read_rows(iter(sheet.values), sheet.title if len(workbook.worksheets) > 1 else ""))
        return deduplicate_papers(papers)
    finally:
        workbook.close()


def write_ris(path: str | Path, papers: list[Paper], include_pdf: bool = False) -> None:
    families = [document_type_family(paper.document_type) for paper in papers]
    unsupported = list(dict.fromkeys(paper.document_type for paper, family in zip(papers, families) if family not in _RIS_TYPES))
    if unsupported:
        raise ValueError(
            f"以下文献类型暂不支持准确的 Zotero 导出：{'、'.join(unsupported)}。"
            "请使用 Excel 或 CSV 保存这些结果。"
        )
    with open(path, "w", encoding="utf-8-sig", newline="\n") as stream:
        def emit(tag: str, value: str) -> None:
            if value:
                stream.write(f"{tag}  - {' '.join(str(value).split())}\n")

        for paper, family in zip(papers, families):
            kind = _RIS_TYPES[family]
            thesis = kind == "THES"
            emit("TY", kind)
            emit("TI", paper.title)
            for author in split_authors(paper.authors):
                emit("AU", author)
            source_tag = {"THES": "PB", "JOUR": "JO", "CONF": "JO", "NEWS": "T2", "CHAP": "T2", "GEN": "JO"}.get(kind)
            if source_tag:
                emit(source_tag, paper.source)
            elif paper.source:
                emit("N1", f"来源：{paper.source}")
            emit("PY", paper.publication_date[:4] if re.match(r"^\d{4}", paper.publication_date) else "")
            emit("DA", paper.publication_date)
            emit("DO", normalize_doi(paper.doi))
            emit("UR", paper.detail_url)
            emit("AB", paper.abstract)
            for keyword in split_values(paper.paper_keywords):
                emit("KW", keyword)
            emit("VL", paper.volume)
            emit("IS", paper.issue)
            if paper.pages:
                page_range = re.split(r"[-–]", paper.pages, maxsplit=1)
                emit("SP", page_range[0])
                if len(page_range) > 1:
                    emit("EP", page_range[1])
            emit("M3", paper.document_type if thesis else "")
            if paper.document_type:
                emit("N1", f"文献类型：{paper.document_type}")
            for query in paper.queries:
                if query.strip():
                    emit("N1", f"命中检索项：{query}")
            if include_pdf and paper.pdf_path and Path(paper.pdf_path).is_file():
                emit("L1", Path(paper.pdf_path).resolve().as_uri())
            for label, value in (("作者单位", paper.institutions), ("基金", paper.funds), ("分类号", paper.classification)):
                if value:
                    emit("N1", f"{label}：{value}")
            stream.write("ER  - \n\n")


def save_papers(path: str | Path, papers: list[Paper], include_pdf: bool = False) -> None:
    from .exporter import _write_atomically

    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        workbook = paper_workbook(papers)
        try:
            _write_atomically(str(path), workbook.save)
        finally:
            workbook.close()
    elif path.suffix.lower() == ".csv":
        _write_atomically(str(path), lambda temporary: write_csv(temporary, papers))
    elif path.suffix.lower() == ".ris":
        _write_atomically(str(path), lambda temporary: write_ris(temporary, papers, include_pdf))
    else:
        raise ValueError("保存格式应为 XLSX、CSV 或 RIS。")
